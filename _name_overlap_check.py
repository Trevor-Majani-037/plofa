"""Compare player names between PLAYER STATS and PLAYERS sheets to find overlaps."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

import unicodedata
import re
import openpyxl

def normalise(name: str) -> str:
    """Strip accents, lowercase, remove non-alphanumeric, collapse whitespace."""
    if not name:
        return ""
    # Strip accents
    nfkd = unicodedata.normalize("NFKD", name)
    stripped = "".join(c for c in nfkd if not unicodedata.combining(c))
    # Lowercase
    stripped = stripped.lower()
    # Remove non-alphanumeric (keep spaces for now)
    stripped = re.sub(r"[^a-z0-9\s]", "", stripped)
    # Collapse whitespace
    stripped = re.sub(r"\s+", " ", stripped).strip()
    return stripped

# ── 1. Read PLAYER STATS from the old workbook ──
old_path = r"D:\TOLAND FOOTBALL FEDERATION\PLOFA-2025-2026.COM\PLOFA 2025-2026 II.xlsx"
print(f"=== Reading PLAYER STATS from: {old_path} ===")
wb_old = openpyxl.load_workbook(old_path, read_only=True, data_only=True)
print(f"Sheet names: {wb_old.sheetnames}")

ws_stats = None
for sn in wb_old.sheetnames:
    if "PLAYER" in sn.upper() and "STAT" in sn.upper():
        ws_stats = wb_old[sn]
        print(f"Using sheet: {sn}")
        break
if ws_stats is None:
    # fallback: try first sheet
    ws_stats = wb_old[wb_old.sheetnames[0]]
    print(f"Fallback to first sheet: {ws_stats.title}")

# Print header row to find column
print("\nHeader row (row 1):")
header_old = []
for i, cell in enumerate(next(ws_stats.iter_rows(min_row=1, max_row=1, values_only=False))):
    header_old.append(str(cell.value) if cell.value else "")
    print(f"  col {i}: {cell.value!r}")

# Column index 4 (0-indexed: 3) for player names
NAME_COL_OLD = 3
print(f"\nUsing column index {NAME_COL_OLD} (header: {header_old[NAME_COL_OLD]!r})")

old_names_raw = []
old_names_norm = []
count = 0
for row in ws_stats.iter_rows(min_row=2, values_only=True):
    if count >= 10:
        break
    val = row[NAME_COL_OLD] if NAME_COL_OLD < len(row) else None
    if val is None:
        continue
    raw = str(val).strip()
    if not raw:
        continue
    nrm = normalise(raw)
    old_names_raw.append(raw)
    old_names_norm.append(nrm)
    count += 1

print(f"\n--- PLAYER STATS: first {len(old_names_raw)} names ---")
for i, (r, n) in enumerate(zip(old_names_raw, old_names_norm)):
    print(f"  {i+1:2d}. RAW: {r!r:40s}  NORM: {n!r}")

# ── 2. Read PLAYERS from the new workbook ──
new_path = r"C:\Users\Trevor Majani\Downloads\plofa_checkpoint6\plofa\PLOFA-2026-2027.xlsx"
print(f"\n=== Reading PLAYERS from: {new_path} ===")
wb_new = openpyxl.load_workbook(new_path, read_only=True, data_only=True)
print(f"Sheet names: {wb_new.sheetnames}")

ws_players = None
for sn in wb_new.sheetnames:
    if sn.upper() == "PLAYERS":
        ws_players = wb_new[sn]
        print(f"Using sheet: {sn}")
        break
if ws_players is None:
    print("WARNING: PLAYERS sheet not found, trying first sheet")
    ws_players = wb_new[wb_new.sheetnames[0]]

print("\nHeader row (row 1):")
header_new = []
name_col_new = None
for i, cell in enumerate(next(ws_players.iter_rows(min_row=1, max_row=1, values_only=False))):
    header_new.append(str(cell.value) if cell.value else "")
    print(f"  col {i}: {cell.value!r}")
    if cell.value and "PLAYER" in str(cell.value).upper() and "NAME" in str(cell.value).upper():
        name_col_new = i

if name_col_new is None:
    # try just "PLAYER"
    for i, h in enumerate(header_new):
        if "PLAYER" in h.upper():
            name_col_new = i
            break

if name_col_new is None:
    print("ERROR: Could not find PLAYER NAME column!")
else:
    print(f"\nUsing column index {name_col_new} (header: {header_new[name_col_new]!r})")

    new_names_raw = []
    new_names_norm = []
    count = 0
    for row in ws_players.iter_rows(min_row=2, values_only=True):
        if count >= 10:
            break
        val = row[name_col_new] if name_col_new < len(row) else None
        if val is None:
            continue
        raw = str(val).strip()
        if not raw:
            continue
        nrm = normalise(raw)
        new_names_raw.append(raw)
        new_names_norm.append(nrm)
        count += 1

    print(f"\n--- PLAYERS: first {len(new_names_raw)} names ---")
    for i, (r, n) in enumerate(zip(new_names_raw, new_names_norm)):
        print(f"  {i+1:2d}. RAW: {r!r:40s}  NORM: {n!r}")

# ── 3. Full overlap check ──
print("\n=== FULL OVERLAP CHECK ===")

# Re-read ALL names from both sheets
old_all_raw = []
old_all_norm = []
for row in ws_stats.iter_rows(min_row=2, values_only=True):
    val = row[NAME_COL_OLD] if NAME_COL_OLD < len(row) else None
    if val is None:
        continue
    raw = str(val).strip()
    if not raw:
        continue
    old_all_raw.append(raw)
    old_all_norm.append(normalise(raw))

new_all_raw = []
new_all_norm = []
for row in ws_players.iter_rows(min_row=2, values_only=True):
    val = row[name_col_new] if name_col_new < len(row) else None
    if val is None:
        continue
    raw = str(val).strip()
    if not raw:
        continue
    new_all_raw.append(raw)
    new_all_norm.append(normalise(raw))

print(f"PLAYER STATS total names: {len(old_all_raw)}")
print(f"PLAYERS total names: {len(new_all_raw)}")

old_set = set(old_all_norm)
new_set = set(new_all_norm)
overlap = old_set & new_set
print(f"Exact normalised matches: {len(overlap)}")
print(f"  PLAYER STATS only: {len(old_set - new_set)}")
print(f"  PLAYERS only: {len(new_set - old_set)}")

if overlap:
    print("\nMatched names (normalised):")
    for n in sorted(overlap):
        # find raw versions
        old_raws = [old_all_raw[i] for i, x in enumerate(old_all_norm) if x == n]
        new_raws = [new_all_raw[i] for i, x in enumerate(new_all_norm) if x == n]
        print(f"  {n!r}")
        print(f"    old raw: {old_raws}")
        print(f"    new raw: {new_raws}")

# Show some near-misses
print("\n--- Near-miss analysis (names in old not in new, similarity check) ---")
import difflib
new_norm_list = list(new_all_norm)
for oname in sorted(old_set - new_set)[:20]:
    matches = difflib.get_close_matches(oname, new_norm_list, n=3, cutoff=0.6)
    if matches:
        print(f"  OLD: {oname!r}  ->  possible matches: {matches}")

wb_old.close()
wb_new.close()
print("\nDone.")
