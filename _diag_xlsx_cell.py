"""Identify PLAYERS!R207 -- which column is R, and whose row is 207?

`_diag_xlsx_diff.py` found exactly one differing cell in the real roster:
PLAYERS!R207 went 44 -> 43. Before anything else, find out what column R means
and which player row 207 is. Then the next step is to find the code that
DECREMENTS it, which is a different search from "what reads it".
"""
import sys
from pathlib import Path

import openpyxl


def main():
    path = Path(__file__).resolve().parent / "PLOFA-2026-2027.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["PLAYERS"]
    print(f"PLAYERS: {ws.max_row} rows x {ws.max_column} cols\n")

    print("header row 1, cols A..T:")
    for col in range(1, min(ws.max_column, 20) + 1):
        letter = openpyxl.utils.get_column_letter(col)
        print(f"  {letter:<3} {ws.cell(row=1, column=col).value!r}")

    print(f"\nrow 207 (the changed row):")
    for col in range(1, min(ws.max_column, 20) + 1):
        letter = openpyxl.utils.get_column_letter(col)
        print(f"  {letter:<3} {ws.cell(row=207, column=col).value!r}")

    # Column R across every row, to see whether 43/44 looks like a counter
    # that varies per player (a stat) or is uniform (a config/global value).
    print("\ncolumn R values, first 30 non-empty:")
    seen = 0
    for row in range(2, ws.max_row + 1):
        v = ws.cell(row=row, column=18).value
        if v is not None:
            print(f"  R{row} = {v!r}")
            seen += 1
            if seen >= 30:
                break
    wb.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
